from __future__ import annotations

import io
import logging
import traceback
from typing import Any, Dict, List, Optional
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

# PDF report generation (charts + narrative)
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import pandas as pd
from fastapi import Depends, FastAPI, HTTPException, Query, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field, validator
from typing import List, Optional
from sqlalchemy.orm import Session

import prediction
import forecasting
from chatbot import chatbot

from database import Base, engine, get_db
from models import AppUser, RecommendationRun, ChatMessage, Feedback, AuditLog
from auth import hash_password, verify_password, create_access_token, get_current_user

# Additional imports for email report functionality
import os
import smtplib
from email.message import EmailMessage
import json
import re

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(
    title="EcoPackAI Backend",
    description="AI-Powered Sustainable Packaging Recommendation System",
    version="1.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # tighten in production
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------- Schemas ----------
class SignupRequest(BaseModel):
    email: str = Field(..., min_length=5)
    password: str = Field(..., min_length=8)
    companyName: str = Field(..., min_length=2)


class LoginRequest(BaseModel):
    email: str
    password: str


class ProductRequest(BaseModel):
    productName: Optional[str] = Field("Generic Product")
    category: Optional[str] = Field("General")
    weightKg: float = Field(..., gt=0)
    fragility: int = Field(..., ge=1, le=10)
    maxBudget: float = Field(..., gt=0)
    shippingDistance: float = Field(500.0, ge=0)
    moistureReq: int = Field(5, ge=0, le=10)
    oxygenSensitivity: int = Field(5, ge=0, le=10)
    preferredBiodegradable: int = Field(0, ge=0, le=1)
    preferredRecyclable: int = Field(0, ge=0, le=1)


class PlannedVolume(BaseModel):
    period: str = Field(..., description="YYYY-MM")
    volumeTons: float = Field(..., gt=0)

    @validator("period")
    def validate_period(cls, value: str) -> str:
        try:
            pd.Period(value, freq="M")
        except Exception as exc:
            raise ValueError("period must be YYYY-MM") from exc
        return value


class ForecastRequest(BaseModel):
    plannedVolumes: List[PlannedVolume]
    simulations: int = Field(400, ge=100, le=2000)


class ChatHistoryMessage(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    question: str
    history: List[ChatHistoryMessage] = []

# ---------- Material Comparison Schema ----------
class CompareMaterialsRequest(BaseModel):
    """
    Schema for comparing specific materials based on user-provided product parameters.
    The client must provide the same fields as used for recommendations along with
    a list of material names to compare.  Up to 10 materials may be compared
    simultaneously.  Unknown materials (i.e. names not present in the training
    dataset) will trigger a call to the OpenAI API to estimate cost and CO₂.
    """
    productName: Optional[str] = Field("Generic Product")
    category: str = Field(..., min_length=1)
    weightKg: float = Field(..., gt=0)
    fragility: int = Field(..., ge=1, le=10)
    maxBudget: float = Field(..., gt=0)
    shippingDistance: float = Field(..., ge=0)
    moistureReq: int = Field(5, ge=0, le=10)
    oxygenSensitivity: int = Field(5, ge=0, le=10)
    preferredBiodegradable: int = Field(0, ge=0, le=1)
    preferredRecyclable: int = Field(0, ge=0, le=1)
    materials: List[str] = Field(..., min_items=1, max_items=10)


# ---------- Preferences & Feedback Schemas ----------
class PreferencesRequest(BaseModel):
    """Schema for updating user-specific model weighting preferences."""
    co2Weight: Optional[float] = Field(None, ge=0.0, le=1.0)
    costWeight: Optional[float] = Field(None, ge=0.0, le=1.0)
    riskWeight: Optional[float] = Field(None, ge=0.0, le=1.0)

    @validator("riskWeight")
    def validate_sum(cls, v: Optional[float], values: Dict[str, Any]) -> Optional[float]:
        co2 = values.get("co2Weight")
        cost = values.get("costWeight")
        if v is not None and co2 is not None and cost is not None:
            if co2 + cost + v > 1.0:
                raise ValueError("Sum of co2Weight, costWeight and riskWeight must be <= 1")
        return v


class FeedbackRequest(BaseModel):
    """Schema for submitting feedback on a recommendation run."""
    runId: int = Field(..., gt=0)
    rating: int = Field(..., ge=1, le=5)
    materialName: Optional[str] = None
    comment: Optional[str] = None


# ---------- Report Email Schema ----------
class ReportEmailRequest(BaseModel):
    """Schema for requesting a report to be emailed."""
    format: Optional[str] = Field('pdf', pattern="^(pdf|excel)$")


# ---------- Startup ----------
@app.on_event("startup")
async def startup_event() -> None:
    try:
        Base.metadata.create_all(bind=engine)
        logger.info("DB tables ensured (ORM).")

        # Ensure forecasting tables use SAME Postgres engine
        forecasting.init_db(engine=engine)
        logger.info("Forecasting DB initialized.")

        prediction.load_models()
        logger.info("ML models loaded.")
    except Exception as e:
        logger.error("Startup failed: %s", e)
        traceback.print_exc()


# ---------- Auth ----------
@app.post("/auth/signup")
def signup(payload: SignupRequest, db: Session = Depends(get_db)) -> JSONResponse:
    email = payload.email.strip().lower()
    company = payload.companyName.strip()

    if db.query(AppUser).filter(AppUser.email == email).first():
        raise HTTPException(status_code=400, detail="Email already registered.")

    user = AppUser(
        email=email,
        password_hash=hash_password(payload.password),
        company_name=company,
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    token = create_access_token(user.id, user.email)
    return JSONResponse(content={"access_token": token, "email": user.email, "companyName": user.company_name})


@app.post("/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)) -> JSONResponse:
    email = payload.email.strip().lower()
    user = db.query(AppUser).filter(AppUser.email == email).first()

    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid email or password.")

    token = create_access_token(user.id, user.email)
    return JSONResponse(content={"access_token": token, "email": user.email, "companyName": user.company_name})


# ---------- Audit logging helper ----------
def log_audit(db: Session, user: Optional[AppUser], action: str, details: Optional[Dict[str, Any]] = None) -> None:
    try:
        log_entry = AuditLog(
            user_id=user.id if user else None,
            action=action,
            details=details or {},
        )
        db.add(log_entry)
        db.commit()
    except Exception:
        db.rollback()


# ---------- Recommend (Protected + Stored) ----------
@app.post("/recommend")
def recommend_materials(
    product: ProductRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    try:
        # Load user-specific weighting preferences if provided.
        weights: Optional[Dict[str, float]] = None
        if getattr(current_user, "weight_co2", None) is not None or getattr(current_user, "weight_cost", None) is not None:
            weights = {}
            if current_user.weight_co2 is not None:
                weights["co2"] = current_user.weight_co2
            if current_user.weight_cost is not None:
                weights["cost"] = current_user.weight_cost

        recs = prediction.recommend(product.dict(), weights=weights)
        top = recs[0] if recs else {}

        run = RecommendationRun(
            user_id=current_user.id,
            product_name=product.productName,
            category=product.category,
            weight_kg=product.weightKg,
            fragility=product.fragility,
            max_budget=product.maxBudget,
            shipping_distance=product.shippingDistance,
            moisture_req=product.moistureReq,
            oxygen_sensitivity=product.oxygenSensitivity,
            preferred_biodegradable=product.preferredBiodegradable,
            preferred_recyclable=product.preferredRecyclable,
            top_material_id=str(top.get("materialId")) if top else None,
            top_material_name=str(top.get("materialName")) if top else None,
            top_pred_cost=float(top.get("predictedCost", 0.0)) if top else None,
            top_pred_co2=float(top.get("predictedCO2", 0.0)) if top else None,
            top_score=float(top.get("suitabilityScore", top.get("rankingScore", 0.0))) if top else None,
            recommendations_json=recs,
        )
        db.add(run)
        db.commit()

        try:
            if top:
                forecasting.log_recommendation_event(
                    engine=engine,
                    event_ts=None,
                    product_category=product.category,
                    product_id=None,
                    material_id=str(top.get("materialId") or top.get("materialName") or "unknown"),
                    material_name=str(top.get("materialName") or "unknown"),
                    pred_cost_per_unit_usd=float(top.get("predictedCost", 0.0)),
                    pred_co2_per_unit_kg=float(top.get("predictedCO2", 0.0)),
                    volume_units=float(product.weightKg),
                )
        except Exception as e:
            logger.warning("Event logging failed: %s", e)

        try:
            log_audit(db, current_user, "recommend", details={
                "product": product.dict(),
                "numRecommendations": len(recs),
                "topMaterial": top.get("materialName") if top else None,
            })
        except Exception:
            pass

        return JSONResponse(content={"recommendations": recs, "runId": run.id})
    except Exception as e:
        logger.error("CRITICAL ERROR in /recommend: %s", e)
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Prediction Failed: {str(e)}")


# ---------- Preferences (Protected) ----------
@app.post("/user/preferences")
def update_preferences(
    payload: PreferencesRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    user = current_user
    co2 = payload.co2Weight
    cost = payload.costWeight
    risk = payload.riskWeight

    current_co2 = user.weight_co2 if user.weight_co2 is not None else 0.5
    current_cost = user.weight_cost if user.weight_cost is not None else 0.35
    current_risk = user.weight_risk if user.weight_risk is not None else 0.15

    new_co2 = co2 if co2 is not None else current_co2
    new_cost = cost if cost is not None else current_cost
    new_risk = risk if risk is not None else current_risk

    total = new_co2 + new_cost + new_risk
    if total > 1.0:
        new_co2 /= total
        new_cost /= total
        new_risk /= total

    user.weight_co2 = new_co2
    user.weight_cost = new_cost
    user.weight_risk = new_risk
    db.add(user)
    db.commit()

    log_audit(db, user, "update_preferences", details={
        "co2Weight": new_co2,
        "costWeight": new_cost,
        "riskWeight": new_risk,
    })

    return JSONResponse(content={"message": "Preferences updated", "co2Weight": new_co2, "costWeight": new_cost, "riskWeight": new_risk})


# ---------- Feedback (Protected) ----------
@app.post("/feedback")
def submit_feedback(
    payload: FeedbackRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    run = db.query(RecommendationRun).filter(
        RecommendationRun.id == payload.runId,
        RecommendationRun.user_id == current_user.id,
    ).first()
    if not run:
        raise HTTPException(status_code=404, detail="Recommendation run not found")

    fb = Feedback(
        run_id=run.id,
        material_name=payload.materialName if payload.materialName else run.top_material_name,
        rating=payload.rating,
        comment=payload.comment,
    )
    db.add(fb)
    db.commit()

    log_audit(db, current_user, "submit_feedback", details={
        "runId": run.id,
        "materialName": fb.material_name,
        "rating": fb.rating,
    })

    return JSONResponse(content={"message": "Feedback submitted"})


# ---------- Trend Analysis (Protected) ----------
@app.get("/trend-data")
def trend_data(
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    runs = (
        db.query(RecommendationRun)
        .filter(RecommendationRun.user_id == current_user.id)
        .order_by(RecommendationRun.created_at.asc())
        .all()
    )

    import pandas as pd

    if not runs:
        return JSONResponse(content={
            "co2Trend": {"labels": [], "values": []},
            "costTrend": {"labels": [], "values": []},
            "usageTrend": {"labels": [], "values": []},
        })

    rows: List[Dict[str, Any]] = []
    for r in runs:
        rows.append({
            "created_at": r.created_at,
            "top_pred_cost": r.top_pred_cost,
            "top_pred_co2": r.top_pred_co2,
            "category": r.category,
            "material": r.top_material_name,
            "preferred_biodegradable": r.preferred_biodegradable,
            "preferred_recyclable": r.preferred_recyclable,
        })
    df = pd.DataFrame(rows)

    df = df.dropna(subset=["top_pred_cost", "top_pred_co2"])
    if df.empty:
        return JSONResponse(content={
            "co2Trend": {"labels": [], "values": []},
            "costTrend": {"labels": [], "values": []},
            "usageTrend": {"labels": [], "values": []},
            "categorySummary": {"categories": [], "runs": [], "avgCost": [], "avgCO2": []},
            "preferences": {"biodegradable": None, "recyclable": None},
        })

    worst_cost = float(pd.to_numeric(df["top_pred_cost"], errors="coerce").dropna().max())
    worst_co2 = float(pd.to_numeric(df["top_pred_co2"], errors="coerce").dropna().max())
    df["month"] = df["created_at"].dt.to_period("M").astype(str)
    monthly = df.groupby("month").agg({
        "top_pred_cost": "mean",
        "top_pred_co2": "mean",
    }).reset_index()
    monthly["co2_reduction"] = (1 - (monthly["top_pred_co2"] / worst_co2)).fillna(0) * 100
    monthly["cost_savings"] = ((worst_cost - monthly["top_pred_cost"]).fillna(0)) * 1.0
    usage = df["material"].fillna("Unknown").astype(str).value_counts(normalize=True) * 100

    cat_df = pd.DataFrame(rows)
    if "category" not in cat_df.columns:
        cat_df["category"] = None
    if "preferred_biodegradable" not in cat_df.columns:
        cat_df["preferred_biodegradable"] = None
    if "preferred_recyclable" not in cat_df.columns:
        cat_df["preferred_recyclable"] = None
    cat_df["category"] = cat_df["category"].fillna("Unknown")

    try:
        cat_df_numeric = cat_df.copy()
        cat_df_numeric["top_pred_cost"] = pd.to_numeric(cat_df_numeric["top_pred_cost"], errors="coerce")
        cat_df_numeric["top_pred_co2"] = pd.to_numeric(cat_df_numeric["top_pred_co2"], errors="coerce")
        category_summary = (
            cat_df_numeric.groupby("category")
            .agg(
                runs=("category", "count"),
                avg_cost=("top_pred_cost", "mean"),
                avg_co2=("top_pred_co2", "mean"),
            )
            .reset_index()
            .sort_values("runs", ascending=False)
        )
        cat_names = category_summary["category"].astype(str).tolist()
        cat_runs = category_summary["runs"].astype(int).tolist()
        cat_avg_cost = category_summary["avg_cost"].round(2).fillna(0).tolist()
        cat_avg_co2 = category_summary["avg_co2"].round(2).fillna(0).tolist()
    except Exception:
        cat_names, cat_runs, cat_avg_cost, cat_avg_co2 = [], [], [], []

    try:
        bio_flags = pd.to_numeric(cat_df["preferred_biodegradable"], errors="coerce").fillna(0)
        rec_flags = pd.to_numeric(cat_df["preferred_recyclable"], errors="coerce").fillna(0)
        bio_rate = float(bio_flags.mean() * 100.0) if len(bio_flags) > 0 else None
        rec_rate = float(rec_flags.mean() * 100.0) if len(rec_flags) > 0 else None
    except Exception:
        bio_rate, rec_rate = None, None

    try:
        total_runs = int(df.shape[0])
        avg_cost = float(pd.to_numeric(df["top_pred_cost"], errors="coerce").mean()) if total_runs > 0 else None
        avg_co2 = float(pd.to_numeric(df["top_pred_co2"], errors="coerce").mean()) if total_runs > 0 else None
    except Exception:
        total_runs, avg_cost, avg_co2 = 0, None, None

    return JSONResponse(content={
        "co2Trend": {
            "labels": monthly["month"].tolist(),
            "values": monthly["co2_reduction"].round(2).tolist(),
        },
        "costTrend": {
            "labels": monthly["month"].tolist(),
            "values": monthly["cost_savings"].round(2).tolist(),
        },
        "usageTrend": {
            "labels": usage.index.astype(str).tolist(),
            "values": usage.round(2).tolist(),
        },
        "categorySummary": {
            "categories": cat_names,
            "runs": cat_runs,
            "avgCost": cat_avg_cost,
            "avgCO2": cat_avg_co2,
        },
        "preferences": {
            "biodegradable": None if bio_rate is None else round(bio_rate, 2),
            "recyclable": None if rec_rate is None else round(rec_rate, 2),
        },
        "summary": {
            "totalRuns": total_runs,
            "avgCost": None if avg_cost is None or pd.isna(avg_cost) else round(avg_cost, 2),
            "avgCO2": None if avg_co2 is None or pd.isna(avg_co2) else round(avg_co2, 2),
        },
    })


# ---------- Sustainability Score (Protected) ----------
@app.get("/sustainability-score")
def sustainability_score(
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    runs = (
        db.query(RecommendationRun)
        .filter(RecommendationRun.user_id == current_user.id)
        .order_by(RecommendationRun.created_at.desc())
        .all()
    )

    import pandas as pd

    if not runs:
        current_user.last_sustainability_score = 0.0
        db.commit()
        return JSONResponse(content={"score": 0.0})

    rows: List[Dict[str, Any]] = []
    for r in runs:
        rows.append({
            "max_budget": r.max_budget,
            "top_pred_cost": r.top_pred_cost,
            "top_pred_co2": r.top_pred_co2,
            "preferred_bio": r.preferred_biodegradable or 0,
        })
    df = pd.DataFrame(rows)
    df["max_budget"] = pd.to_numeric(df["max_budget"], errors="coerce")
    df["top_pred_cost"] = pd.to_numeric(df["top_pred_cost"], errors="coerce")
    df["top_pred_co2"] = pd.to_numeric(df["top_pred_co2"], errors="coerce")

    df = df.dropna()
    if df.empty:
        current_user.last_sustainability_score = 0.0
        db.commit()
        return JSONResponse(content={"score": 0.0})

    worst_co2 = float(df["top_pred_co2"].max())
    co2_reduction = (1 - (df["top_pred_co2"] / worst_co2)).mean() if worst_co2 > 0 else 0
    biodegradable_ratio = df["preferred_bio"].mean()
    budget_eff = ((df["max_budget"] - df["top_pred_cost"]) / df["max_budget"]).clip(lower=0).mean()

    score = (0.4 * co2_reduction) + (0.3 * biodegradable_ratio) + (0.3 * budget_eff)
    score = float(max(0.0, min(1.0, score))) * 100.0

    current_user.last_sustainability_score = score
    db.add(current_user)
    db.commit()

    log_audit(db, current_user, "compute_sustainability_score", details={"score": score})

    return JSONResponse(content={"score": round(score, 2)})


# ---------- Dashboard (Protected) ----------
@app.get("/dashboard-data")
def dashboard_data(
    current_user: AppUser = Depends(get_current_user),
) -> JSONResponse:
    try:
        metrics = prediction.compute_dashboard_metrics()
        return JSONResponse(content=metrics)
    except Exception as e:
        logger.error("Dashboard error: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


# ---------- Reports (Protected) ----------
@app.get("/report")
def download_report(
    format: str = Query(..., pattern="^(pdf|excel)$"),
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    runs = (
        db.query(RecommendationRun)
        .filter(RecommendationRun.user_id == current_user.id)
        .order_by(RecommendationRun.created_at.desc())
        .limit(500)
        .all()
    )

    if format == "excel":
        has_excel_engine: bool = False
        try:
            import xlsxwriter  # type: ignore  # noqa: F401
            has_excel_engine = True
        except Exception:
            try:
                import openpyxl  # type: ignore  # noqa: F401
                has_excel_engine = True
            except Exception:
                has_excel_engine = False
        data = _generate_excel_report_from_runs(runs, current_user)
        if has_excel_engine:
            filename = "ecopackai_report.xlsx"
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            filename = "ecopackai_report.csv"
            media_type = "text/csv"
    else:
        data = _generate_pdf_report_from_runs(runs, current_user)
        filename = "ecopackai_report.pdf"
        media_type = "application/pdf"

    return StreamingResponse(
        io.BytesIO(data),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ---------- Email Report (Protected) ----------
@app.post("/report/email")
def email_report(
    payload: ReportEmailRequest,
    background_tasks: BackgroundTasks,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    runs = (
        db.query(RecommendationRun)
        .filter(RecommendationRun.user_id == current_user.id)
        .order_by(RecommendationRun.created_at.desc())
        .limit(500)
        .all()
    )
    fmt = (payload.format or 'pdf').lower()
    if fmt == 'excel':
        has_excel_engine = False
        try:
            import xlsxwriter  # noqa: F401
            has_excel_engine = True
        except Exception:
            try:
                import openpyxl  # noqa: F401
                has_excel_engine = True
            except Exception:
                has_excel_engine = False
        report_data = _generate_excel_report_from_runs(runs, current_user)
        if has_excel_engine:
            filename = "ecopackai_report.xlsx"
            mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            filename = "ecopackai_report.csv"
            mime_type = "text/csv"
    else:
        report_data = _generate_pdf_report_from_runs(runs, current_user)
        filename = "ecopackai_report.pdf"
        mime_type = "application/pdf"

    background_tasks.add_task(
        _send_report_email,
        to_email=current_user.email,
        subject="Your EcoPackAI sustainability report",
        body="Please find attached your EcoPackAI sustainability report.",
        attachments=[(filename, report_data, mime_type)],
    )

    try:
        log_audit(db, current_user, "email_report", details={"format": fmt})
    except Exception:
        pass

    return JSONResponse(content={"message": "Report is being emailed"})


def _send_report_email(
    to_email: str,
    subject: str,
    body: str,
    attachments: List[tuple],
) -> None:
    try:
        host = os.getenv("SMTP_HOST", "smtp.gmail.com")
        port = int(os.getenv("SMTP_PORT", "587"))
        user = os.getenv("SMTP_USER")
        password = os.getenv("SMTP_PASS")
        from_addr = os.getenv("SMTP_FROM", user or "")
        use_tls = os.getenv("SMTP_USE_TLS", "true").lower() == "true"
        use_ssl = os.getenv("SMTP_USE_SSL", "false").lower() == "true"

        if not user or not password:
            logger.error("SMTP credentials are not configured; email not sent")
            return

        msg = EmailMessage()
        msg["Subject"] = subject
        msg["From"] = from_addr
        msg["To"] = to_email
        msg.set_content(body)

        for fname, data, mime in attachments:
            maintype, _, subtype = mime.partition("/")
            msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=fname)

        if use_ssl:
            with smtplib.SMTP_SSL(host, port) as server:
                server.login(user, password)
                server.send_message(msg)
        else:
            with smtplib.SMTP(host, port) as server:
                if use_tls:
                    server.starttls()
                server.login(user, password)
                server.send_message(msg)
        logger.info("Report email sent to %s", to_email)
    except Exception as e:
        logger.error("Failed to send report email: %s", e)


def _runs_to_dataframe(runs: List[RecommendationRun]) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for r in runs:
        rows.append({
            "created_at": r.created_at,
            "product_name": r.product_name,
            "category": r.category,
            "weight_kg": r.weight_kg,
            "fragility": r.fragility,
            "max_budget": r.max_budget,
            "shipping_distance": r.shipping_distance,
            "top_material_name": r.top_material_name,
            "top_pred_cost": r.top_pred_cost,
            "top_pred_co2": r.top_pred_co2,
            "top_score": r.top_score,
            "preferred_biodegradable": getattr(r, "preferred_biodegradable", None),
            "preferred_recyclable": getattr(r, "preferred_recyclable", None),
        })
    return pd.DataFrame(rows)


def _compute_business_metrics(df: pd.DataFrame) -> Dict[str, Any]:
    if df.empty:
        return {
            "summary": {},
            "monthly": pd.DataFrame(columns=["month", "avg_co2", "avg_cost", "co2_reduction_pct", "cost_savings_pct"]),
            "usage": pd.Series(dtype=float),
        }

    dfx = df.copy()
    dfx["created_at"] = pd.to_datetime(dfx["created_at"], errors="coerce")
    dfx["top_pred_cost"] = pd.to_numeric(dfx["top_pred_cost"], errors="coerce")
    dfx["top_pred_co2"] = pd.to_numeric(dfx["top_pred_co2"], errors="coerce")
    dfx["top_material_name"] = dfx["top_material_name"].fillna("Unknown")

    valid_cost = dfx["top_pred_cost"].dropna()
    valid_co2 = dfx["top_pred_co2"].dropna()
    baseline_cost = float(valid_cost.quantile(0.9)) if not valid_cost.empty else None
    baseline_co2 = float(valid_co2.quantile(0.9)) if not valid_co2.empty else None

    dfx["month"] = dfx["created_at"].dt.to_period("M").astype(str)
    monthly = (
        dfx.groupby("month", dropna=True)
        .agg(avg_cost=("top_pred_cost", "mean"), avg_co2=("top_pred_co2", "mean"), runs=("month", "count"))
        .reset_index()
        .sort_values("month")
    )

    if baseline_co2 and baseline_co2 > 0:
        monthly["co2_reduction_pct"] = ((baseline_co2 - monthly["avg_co2"]) / baseline_co2 * 100).clip(lower=-100, upper=100)
    else:
        monthly["co2_reduction_pct"] = 0.0

    if baseline_cost and baseline_cost > 0:
        monthly["cost_savings_pct"] = ((baseline_cost - monthly["avg_cost"]) / baseline_cost * 100).clip(lower=-100, upper=100)
    else:
        monthly["cost_savings_pct"] = 0.0

    usage = dfx["top_material_name"].value_counts(normalize=True) * 100

    try:
        cat_group = dfx.copy()
        cat_group["category"] = cat_group["category"].fillna("Unknown")
        category_summary = (
            cat_group.groupby("category", dropna=False)
            .agg(runs=("category", "count"), avg_cost=("top_pred_cost", "mean"), avg_co2=("top_pred_co2", "mean"))
            .reset_index()
            .sort_values("runs", ascending=False)
        )
    except Exception:
        category_summary = pd.DataFrame(columns=["category", "runs", "avg_cost", "avg_co2"])

    try:
        bio_flags = pd.to_numeric(dfx.get("preferred_biodegradable"), errors="coerce").dropna()
        rec_flags = pd.to_numeric(dfx.get("preferred_recyclable"), errors="coerce").dropna()
        bio_pref_pct = float(bio_flags.mean() * 100.0) if not bio_flags.empty else None
        rec_pref_pct = float(rec_flags.mean() * 100.0) if not rec_flags.empty else None
    except Exception:
        bio_pref_pct = None
        rec_pref_pct = None

    summary = {
        "total_runs": int(len(dfx)),
        "baseline_cost": baseline_cost,
        "baseline_co2": baseline_co2,
        "avg_cost": float(valid_cost.mean()) if not valid_cost.empty else None,
        "avg_co2": float(valid_co2.mean()) if not valid_co2.empty else None,
        "most_common_material": str(usage.index[0]) if not usage.empty else None,
        "most_common_material_pct": float(usage.iloc[0]) if not usage.empty else None,
        "bio_pref_pct": bio_pref_pct,
        "recycle_pref_pct": rec_pref_pct,
    }

    return {
        "summary": summary,
        "monthly": monthly,
        "usage": usage,
        "categories": category_summary,
    }


def _generate_excel_report_from_runs(runs: List[RecommendationRun], user: AppUser) -> bytes:
    df = _runs_to_dataframe(runs)
    metrics = _compute_business_metrics(df)
    monthly: pd.DataFrame = metrics["monthly"]
    usage: pd.Series = metrics["usage"]

    buffer = io.BytesIO()

    try:
        import xlsxwriter  # type: ignore

        wb = xlsxwriter.Workbook(buffer, {"in_memory": True})
        fmt_title = wb.add_format({"bold": True, "font_size": 14})
        fmt_hdr = wb.add_format({"bold": True, "bg_color": "#F2F2F2", "border": 1})
        fmt_cell = wb.add_format({"border": 1})
        fmt_pct = wb.add_format({"num_format": "0.0%", "border": 1})
        fmt_num = wb.add_format({"num_format": "0.000", "border": 1})

        ws = wb.add_worksheet("Executive_Summary")
        ws.write(0, 0, "EcoPackAI Sustainability & Cost Report", fmt_title)
        ws.write(2, 0, "Company")
        ws.write(2, 1, user.company_name or "")
        ws.write(3, 0, "User")
        ws.write(3, 1, user.email or "")
        ws.write(4, 0, "Generated")
        ws.write(4, 1, datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"))

        summ = metrics["summary"]
        ws.write(6, 0, "Total recommendation runs")
        ws.write(6, 1, summ.get("total_runs") or 0)
        ws.write(7, 0, "Average predicted cost (top-1)")
        ws.write(7, 1, summ.get("avg_cost") or 0)
        ws.write(8, 0, "Average predicted CO₂ (top-1)")
        ws.write(8, 1, summ.get("avg_co2") or 0)
        ws.write(9, 0, "Most selected material")
        ws.write(9, 1, summ.get("most_common_material") or "")
        ws.write(10, 0, "Most selected material share")
        pct = (summ.get("most_common_material_pct") or 0) / 100.0
        ws.write(10, 1, pct, fmt_pct)
        ws.write(11, 0, "Biodegradable preference rate")
        bio_pct = (summ.get("bio_pref_pct") or 0) / 100.0
        ws.write(11, 1, bio_pct, fmt_pct)
        ws.write(12, 0, "Recyclable preference rate")
        rec_pct = (summ.get("recycle_pref_pct") or 0) / 100.0
        ws.write(12, 1, rec_pct, fmt_pct)

        ws.write(14, 0, "How to read this report", wb.add_format({"bold": True}))
        ws.write(15, 0, "• Cost and CO₂ are model predictions for the best-ranked material per run.")
        ws.write(16, 0, "• Trends compare monthly averages to a baseline (90th percentile of past runs).")
        ws.write(17, 0, "• Use this report for audits, procurement decisions, and stakeholder updates.")
        ws.set_column(0, 0, 40)
        ws.set_column(1, 1, 30)

        ws_runs = wb.add_worksheet("Runs")
        for c, col in enumerate(df.columns):
            ws_runs.write(0, c, col, fmt_hdr)
        for r_i, row in enumerate(df.itertuples(index=False), start=1):
            for c, val in enumerate(row):
                ws_runs.write(r_i, c, "" if val is None else val, fmt_cell)
        ws_runs.freeze_panes(1, 0)
        ws_runs.set_column(0, 0, 22)
        ws_runs.set_column(1, 2, 18)
        ws_runs.set_column(3, 9, 16)

        ws_tr = wb.add_worksheet("Trends")
        trend_cols = ["month", "avg_cost", "avg_co2", "co2_reduction_pct", "cost_savings_pct", "runs"]
        for c, col in enumerate(trend_cols):
            ws_tr.write(0, c, col, fmt_hdr)
        for r_i, row in enumerate(monthly[trend_cols].itertuples(index=False), start=1):
            ws_tr.write(r_i, 0, getattr(row, "month"), fmt_cell)
            ws_tr.write(r_i, 1, float(getattr(row, "avg_cost") or 0), fmt_num)
            ws_tr.write(r_i, 2, float(getattr(row, "avg_co2") or 0), fmt_num)
            ws_tr.write(r_i, 3, float(getattr(row, "co2_reduction_pct") or 0) / 100.0, fmt_pct)
            ws_tr.write(r_i, 4, float(getattr(row, "cost_savings_pct") or 0) / 100.0, fmt_pct)
            ws_tr.write(r_i, 5, int(getattr(row, "runs") or 0), fmt_cell)

        ws_tr.set_column(0, 0, 12)
        ws_tr.set_column(1, 2, 14)
        ws_tr.set_column(3, 4, 18)
        ws_tr.set_column(5, 5, 8)

        chart1 = wb.add_chart({"type": "line"})
        chart1.add_series({
            "name": "CO₂ reduction vs baseline",
            "categories": ["Trends", 1, 0, max(1, len(monthly)), 0],
            "values": ["Trends", 1, 3, max(1, len(monthly)), 3],
        })
        chart1.set_title({"name": "CO₂ Reduction Over Time"})
        chart1.set_y_axis({"num_format": "0%"})
        ws_tr.insert_chart(1, 7, chart1, {"x_scale": 1.2, "y_scale": 1.2})

        chart2 = wb.add_chart({"type": "line"})
        chart2.add_series({
            "name": "Cost savings vs baseline",
            "categories": ["Trends", 1, 0, max(1, len(monthly)), 0],
            "values": ["Trends", 1, 4, max(1, len(monthly)), 4],
        })
        chart2.set_title({"name": "Cost Savings Over Time"})
        chart2.set_y_axis({"num_format": "0%"})
        ws_tr.insert_chart(17, 7, chart2, {"x_scale": 1.2, "y_scale": 1.2})

        ws_use = wb.add_worksheet("Material_Usage")
        ws_use.write(0, 0, "material", fmt_hdr)
        ws_use.write(0, 1, "usage_pct", fmt_hdr)
        for i, (mat, pct_val) in enumerate(usage.items(), start=1):
            ws_use.write(i, 0, str(mat), fmt_cell)
            ws_use.write(i, 1, float(pct_val) / 100.0, fmt_pct)
        ws_use.set_column(0, 0, 32)
        ws_use.set_column(1, 1, 12)

        chart3 = wb.add_chart({"type": "pie"})
        chart3.add_series({
            "name": "Material usage",
            "categories": ["Material_Usage", 1, 0, max(1, len(usage)), 0],
            "values": ["Material_Usage", 1, 1, max(1, len(usage)), 1],
        })
        chart3.set_title({"name": "Material Usage Distribution"})
        ws_use.insert_chart(1, 3, chart3, {"x_scale": 1.2, "y_scale": 1.2})

        categories_df = metrics.get("categories")
        try:
            import pandas as _pd  # type: ignore
        except Exception:
            _pd = None
        if categories_df is not None and getattr(categories_df, 'empty', True) is False:
            ws_cat = wb.add_worksheet("Category_Summary")
            ws_cat.write(0, 0, "Category", fmt_hdr)
            ws_cat.write(0, 1, "Runs", fmt_hdr)
            ws_cat.write(0, 2, "Avg cost (USD/unit)", fmt_hdr)
            ws_cat.write(0, 3, "Avg CO₂ (kg/unit)", fmt_hdr)
            try:
                cat_df_sorted = categories_df.sort_values("runs", ascending=False).reset_index(drop=True)
            except Exception:
                cat_df_sorted = categories_df
            row_idx = 1
            for _, r in cat_df_sorted.iterrows():
                ws_cat.write(row_idx, 0, str(r.get("category") or "Unknown"), fmt_cell)
                ws_cat.write(row_idx, 1, int(r.get("runs") or 0), fmt_cell)
                ws_cat.write(row_idx, 2, float(r.get("avg_cost") or 0), fmt_num)
                ws_cat.write(row_idx, 3, float(r.get("avg_co2") or 0), fmt_num)
                row_idx += 1
            ws_cat.set_column(0, 0, 30)
            ws_cat.set_column(1, 1, 8)
            ws_cat.set_column(2, 3, 18)

        ws_latest = wb.add_worksheet("Latest_Recommendations")
        ws_latest.write(0, 0, "This sheet summarizes the most recent recommendation run.")
        if runs:
            latest = runs[0]
            ws_latest.write(2, 0, "Product")
            ws_latest.write(2, 1, latest.product_name)
            ws_latest.write(3, 0, "Category")
            ws_latest.write(3, 1, latest.category)
            ws_latest.write(5, 0, "Rank", fmt_hdr)
            ws_latest.write(5, 1, "Material", fmt_hdr)
            ws_latest.write(5, 2, "Cost (USD/unit)", fmt_hdr)
            ws_latest.write(5, 3, "CO₂ (kg/unit)", fmt_hdr)
            ws_latest.write(5, 4, "Score", fmt_hdr)
            ws_latest.write(5, 5, "Reason", fmt_hdr)

            try:
                recs = latest.recommendations_json
                if isinstance(recs, str):
                    import json
                    recs_data = json.loads(recs)
                else:
                    recs_data = recs or []
            except Exception:
                recs_data = []

            for i, rec in enumerate(recs_data[:10], start=1):
                name = rec.get("materialName") or rec.get("MaterialName") or rec.get("MaterialType") or "Unknown"
                cost = rec.get("predictedCost") or rec.get("predictedCostUSD") or rec.get("predicted_cost_unit_usd")
                co2 = rec.get("predictedCO2") or rec.get("predictedCO2KG") or rec.get("predicted_co2_unit_kg")
                score = rec.get("suitabilityScore") or rec.get("rankingScore")
                reason = rec.get("explanation") or rec.get("reason") or ""
                ws_latest.write(5 + i, 0, i, fmt_cell)
                ws_latest.write(5 + i, 1, str(name), fmt_cell)
                ws_latest.write(5 + i, 2, float(cost) if cost is not None else "", fmt_num)
                ws_latest.write(5 + i, 3, float(co2) if co2 is not None else "", fmt_num)
                ws_latest.write(5 + i, 4, float(score) if score is not None else "", fmt_cell)
                ws_latest.write(5 + i, 5, str(reason), fmt_cell)
            ws_latest.set_column(0, 0, 6)
            ws_latest.set_column(1, 1, 30)
            ws_latest.set_column(2, 4, 16)
            ws_latest.set_column(5, 5, 42)

        ws_g = wb.add_worksheet("Glossary")
        ws_g.write(0, 0, "Term", fmt_hdr)
        ws_g.write(0, 1, "Plain-language meaning", fmt_hdr)
        glossary = [
            ("Predicted cost", "Estimated packaging cost per unit for a selected material (model output)."),
            ("Predicted CO₂", "Estimated CO₂ emissions per unit for a selected material (model output)."),
            ("Suitability score", "0–100 score combining constraints and preferences (higher is better)."),
            ("Baseline", "Reference level calculated from your historical runs (90th percentile)."),
            ("CO₂ reduction", "How much lower monthly CO₂ is compared to the baseline."),
            ("Cost savings", "How much lower monthly cost is compared to the baseline."),
        ]
        for i, (t, d) in enumerate(glossary, start=1):
            ws_g.write(i, 0, t, fmt_cell)
            ws_g.write(i, 1, d, fmt_cell)
        ws_g.set_column(0, 0, 22)
        ws_g.set_column(1, 1, 88)

        wb.close()
        buffer.seek(0)
        return buffer.getvalue()

    except Exception:
        try:
            import openpyxl  # noqa: F401
            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="runs", index=False)
                monthly.to_excel(writer, sheet_name="trends", index=False)
                usage.reset_index().rename(columns={"index": "material", 0: "usage_pct"}).to_excel(writer, sheet_name="usage", index=False)
            buffer.seek(0)
            return buffer.getvalue()
        except Exception:
            return df.to_csv(index=False).encode("utf-8")


def _generate_pdf_report_from_runs(runs: List[RecommendationRun], user: AppUser) -> bytes:
    df = _runs_to_dataframe(runs)
    metrics = _compute_business_metrics(df)
    monthly: pd.DataFrame = metrics["monthly"]
    usage: pd.Series = metrics["usage"]
    summ = metrics["summary"]

    styles = getSampleStyleSheet()
    story: List[Any] = []

    story.append(Paragraph("EcoPackAI Sustainability & Cost Report", styles["Title"]))
    story.append(Paragraph(f"<b>Company:</b> {user.company_name or ''}<br/><b>User:</b> {user.email or ''}<br/><b>Generated:</b> {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Executive Summary", styles["Heading2"]))
    story.append(Paragraph(
        "This report summarizes historical recommendation activity, estimated environmental impact (CO₂), and unit cost trends. "
        "It is designed to be readable by non-technical stakeholders while still providing enough detail for technical review.",
        styles["Normal"],
    ))
    story.append(Spacer(1, 8))

    key_rows = [
        ["Metric", "Value"],
        ["Total recommendation runs", str(summ.get("total_runs") or 0)],
        ["Average predicted cost (top-1)", f"{(summ.get('avg_cost') or 0):.3f} USD/unit" if summ.get('avg_cost') is not None else "—"],
        ["Average predicted CO₂ (top-1)", f"{(summ.get('avg_co2') or 0):.3f} kg/unit" if summ.get('avg_co2') is not None else "—"],
        ["Most selected material", f"{summ.get('most_common_material') or '—'}"],
        ["Most selected material share", f"{(summ.get('most_common_material_pct') or 0):.1f}%" if summ.get('most_common_material_pct') is not None else "—"],
        ["Biodegradable preference rate", f"{(summ.get('bio_pref_pct') or 0):.1f}%" if summ.get('bio_pref_pct') is not None else "—"],
        ["Recyclable preference rate", f"{(summ.get('recycle_pref_pct') or 0):.1f}%" if summ.get('recycle_pref_pct') is not None else "—"],
    ]
    tbl = Table(key_rows, colWidths=[200, 320])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 14))

    def _fig_to_rl_image(fig, width=520):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", bbox_inches="tight", dpi=150)
        plt.close(fig)
        buf.seek(0)
        img = RLImage(buf)
        scale = width / img.drawWidth
        img.drawWidth = width
        img.drawHeight = img.drawHeight * scale
        return img

    if not monthly.empty:
        story.append(Paragraph("Business Trends", styles["Heading2"]))
        story.append(Paragraph(
            "Trends are calculated as monthly averages compared to a baseline derived from your historical runs (90th percentile). "
            "Positive values indicate improvement (lower CO₂ or lower cost).",
            styles["Normal"],
        ))
        story.append(Spacer(1, 8))

        fig1 = plt.figure(figsize=(7.2, 3.2))
        ax1 = fig1.add_subplot(111)
        ax1.plot(monthly["month"], monthly["co2_reduction_pct"].fillna(0.0))
        ax1.set_title("CO₂ Reduction Over Time (% vs baseline)")
        ax1.set_xlabel("Month")
        ax1.set_ylabel("%")
        ax1.tick_params(axis='x', rotation=45)
        story.append(_fig_to_rl_image(fig1))
        story.append(Spacer(1, 8))

        fig2 = plt.figure(figsize=(7.2, 3.2))
        ax2 = fig2.add_subplot(111)
        ax2.plot(monthly["month"], monthly["cost_savings_pct"].fillna(0.0))
        ax2.set_title("Cost Savings Over Time (% vs baseline)")
        ax2.set_xlabel("Month")
        ax2.set_ylabel("%")
        ax2.tick_params(axis='x', rotation=45)
        story.append(_fig_to_rl_image(fig2))
        story.append(Spacer(1, 14))

    if not usage.empty:
        story.append(Paragraph("Material Usage Distribution", styles["Heading2"]))
        story.append(Paragraph(
            "This chart shows how often each top material was selected across runs. "
            "High concentration may indicate standardized procurement; diversification may indicate experimentation.",
            styles["Normal"],
        ))
        story.append(Spacer(1, 8))
        fig3 = plt.figure(figsize=(7.2, 3.6))
        ax3 = fig3.add_subplot(111)
        ax3.pie(usage.values, labels=usage.index.astype(str), autopct="%1.0f%%")
        ax3.set_title("Top Material Usage")
        story.append(_fig_to_rl_image(fig3))
        story.append(Spacer(1, 14))

    try:
        cat_df = metrics.get("categories")
    except Exception:
        cat_df = None
    if cat_df is not None and not isinstance(cat_df, list) and not getattr(cat_df, 'empty', True):
        story.append(Paragraph("Category Performance Summary", styles["Heading2"]))
        story.append(Paragraph(
            "Average predicted cost and CO₂ impact per product category across your recommendation runs.",
            styles["Normal"],
        ))
        story.append(Spacer(1, 8))
        rows = [["Category", "Runs", "Avg cost (USD/unit)", "Avg CO₂ (kg/unit)"]]
        try:
            top_cats = cat_df.copy().head(10)
            for _, r in top_cats.iterrows():
                rows.append([
                    str(r.get("category", "Unknown")),
                    str(int(r.get("runs") or 0)),
                    f"{float(r.get('avg_cost') or 0):.3f}" if pd.notna(r.get('avg_cost')) else "",
                    f"{float(r.get('avg_co2') or 0):.3f}" if pd.notna(r.get('avg_co2')) else "",
                ])
        except Exception:
            pass
        cat_tbl = Table(rows, colWidths=[150, 60, 150, 150])
        cat_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(cat_tbl)
        story.append(Spacer(1, 14))

    if summ.get("bio_pref_pct") is not None or summ.get("recycle_pref_pct") is not None:
        story.append(Paragraph("Preferences Summary", styles["Heading2"]))
        story.append(Paragraph(
            "Share of recommendation runs where biodegradable and recyclable materials were preferred.",
            styles["Normal"],
        ))
        story.append(Spacer(1, 8))
        pref_rows = [["Preference", "Rate (%)"]]
        bio_rate = summ.get("bio_pref_pct")
        rec_rate = summ.get("recycle_pref_pct")
        pref_rows.append(["Biodegradable", f"{float(bio_rate):.1f}%" if bio_rate is not None else "—"])
        pref_rows.append(["Recyclable", f"{float(rec_rate):.1f}%" if rec_rate is not None else "—"])
        pref_tbl = Table(pref_rows, colWidths=[200, 200])
        pref_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(pref_tbl)
        story.append(Spacer(1, 14))

    if cat_df is not None and not isinstance(cat_df, list) and not getattr(cat_df, 'empty', True):
        try:
            story.append(Paragraph("Category Performance Chart", styles["Heading2"]))
            story.append(Paragraph(
                "Average predicted cost and CO₂ impact for top product categories across your runs.",
                styles["Normal"],
            ))
            story.append(Spacer(1, 8))
            top_cats = cat_df.copy().sort_values("runs", ascending=False).head(6)
            categories = top_cats["category"].astype(str).tolist()
            avg_costs = top_cats["avg_cost"].fillna(0).tolist()
            avg_co2s = top_cats["avg_co2"].fillna(0).tolist()
            fig_bar = plt.figure(figsize=(7.2, 3.6))
            ax_bar = fig_bar.add_subplot(111)
            import numpy as _np
            x = _np.arange(len(categories))
            width = 0.35
            ax_bar.bar(x - width/2, avg_costs, width=width, label='Avg cost (USD/unit)')
            ax_bar.bar(x + width/2, avg_co2s, width=width, label='Avg CO₂ (kg/unit)')
            ax_bar.set_xticks(x)
            ax_bar.set_xticklabels(categories, rotation=45, ha='right')
            ax_bar.set_ylabel("Value")
            ax_bar.set_title("Category cost and CO₂ impact")
            ax_bar.legend()
            story.append(_fig_to_rl_image(fig_bar))
            story.append(Spacer(1, 14))
        except Exception as e:
            logger.warning("Failed to build category bar chart: %s", e)

    story.append(Paragraph("Latest Recommendation Snapshot", styles["Heading2"]))
    if runs:
        latest = runs[0]
        story.append(Paragraph(
            f"<b>Product:</b> {latest.product_name or ''} &nbsp;&nbsp; <b>Category:</b> {latest.category or ''}",
            styles["Normal"],
        ))
        story.append(Spacer(1, 8))
        try:
            recs = latest.recommendations_json
            if isinstance(recs, str):
                import json
                recs_data = json.loads(recs)
            else:
                recs_data = recs or []
        except Exception:
            recs_data = []

        rows = [["Rank", "Material", "Cost (USD/unit)", "CO₂ (kg/unit)", "Score", "Reason"]]
        for i, rec in enumerate(recs_data[:5], start=1):
            name = rec.get("materialName") or rec.get("MaterialName") or rec.get("MaterialType") or "Unknown"
            cost = rec.get("predictedCost") or rec.get("predictedCostUSD") or rec.get("predicted_cost_unit_usd")
            co2 = rec.get("predictedCO2") or rec.get("predictedCO2KG") or rec.get("predicted_co2_unit_kg")
            score = rec.get("suitabilityScore") or rec.get("rankingScore")
            reason = rec.get("explanation") or rec.get("reason") or ""
            rows.append([
                str(i),
                str(name),
                f"{float(cost):.3f}" if cost is not None else "",
                f"{float(co2):.3f}" if co2 is not None else "",
                f"{float(score):.0f}" if score is not None else "",
                str(reason)[:120],
            ])

        rec_tbl = Table(rows, colWidths=[40, 150, 80, 70, 50, 160])
        rec_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(rec_tbl)
    else:
        story.append(Paragraph("No recommendation runs found for this user.", styles["Normal"]))
    story.append(Spacer(1, 14))

    story.append(Paragraph("Glossary", styles["Heading2"]))
    story.append(Paragraph(
        "<b>Predicted cost</b>: estimated unit cost for a selected material (model output).<br/>"
        "<b>Predicted CO₂</b>: estimated emissions per unit (model output).<br/>"
        "<b>Suitability score</b>: 0–100 combined measure of constraint fit + preferences.<br/>"
        "<b>Baseline</b>: a reference derived from historical runs (90th percentile).",
        styles["Normal"],
    ))

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=letter, title="EcoPackAI Report")
    doc.build(story)
    out.seek(0)
    return out.getvalue()

# ---------- Recommendation History (Protected) ----------
@app.get("/recommend/history")
def get_recommendation_history(
    limit: int = Query(20, ge=1, le=500, description="Maximum number of recent runs to return"),
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    runs = (
        db.query(RecommendationRun)
        .filter(RecommendationRun.user_id == current_user.id)
        .order_by(RecommendationRun.created_at.desc())
        .limit(limit)
        .all()
    )

    history: List[Dict[str, Any]] = []
    for r in runs:
        try:
            created_at_iso = r.created_at.isoformat() if r.created_at else None
        except Exception:
            created_at_iso = None
        history.append({
            "id": r.id,
            "createdAt": created_at_iso,
            "productName": r.product_name,
            "category": r.category,
            "weightKg": r.weight_kg,
            "fragility": r.fragility,
            "maxBudget": r.max_budget,
            "shippingDistance": r.shipping_distance,
            "moistureReq": r.moisture_req,
            "oxygenSensitivity": r.oxygen_sensitivity,
            "preferredBiodegradable": r.preferred_biodegradable,
            "preferredRecyclable": r.preferred_recyclable,
            "topMaterialName": r.top_material_name,
            "topScore": r.top_score,
            "topPredCost": r.top_pred_cost,
            "topPredCO2": r.top_pred_co2,
        })

    try:
        log_audit(db, current_user, "list_history", details={"count": len(history)})
    except Exception:
        pass

    return JSONResponse(content={"history": history})


# ---------- Forecast (Protected) ----------
@app.post("/forecast")
def forecast_trends(
    payload: ForecastRequest,
    current_user: AppUser = Depends(get_current_user),
) -> JSONResponse:
    try:
        if payload.plannedVolumes and len(payload.plannedVolumes) > 0:
            planned = [{"period": pv.period, "volumeTons": pv.volumeTons} for pv in payload.plannedVolumes]
            result = forecasting.generate_plan_impact_forecast(
                engine=engine,
                planned=planned,
                freq="M",
                simulations=int(payload.simulations),
            )
            return JSONResponse(content=result)

        result = forecasting.generate_forecast_from_events(
            engine=engine,
            horizon_periods=6,
            freq="M",
            simulations=int(payload.simulations),
        )
        return JSONResponse(content=result)

    except Exception as e:
        logger.error("Forecast error: %s", e)
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))


# ---------- Chat (Protected + Stored) ----------
@app.post("/chat")
def chat_with_assistant(
    payload: ChatRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> JSONResponse:
    if not payload.question.strip():
        raise HTTPException(status_code=400, detail="Empty question")

    try:
        db.add(ChatMessage(user_id=current_user.id, role="user", content=payload.question))
        db.commit()

        q_lower = payload.question.strip().lower()
        history = [m.dict() for m in payload.history]

        import os as _os
        try:
            _os.environ["CHATBOT_CURRENT_USER_EMAIL"] = current_user.email or ""
            _os.environ["CHATBOT_CURRENT_COMPANY"] = current_user.company_name or ""
        except Exception:
            pass

        history_for_bot = history.copy() if history else []
        try:
            sys_content = (
                f"The current user is {current_user.email or 'unknown'} and works at "
                f"{current_user.company_name or 'unknown'}. Answer questions with this context in mind."
            )
            history_for_bot.insert(0, {"role": "system", "content": sys_content})
        except Exception:
            history_for_bot = history

        answered = False
        answer: str = ""
        try:
            if ("why" in q_lower and "recommended" in q_lower and "over" in q_lower) or ("last run" in q_lower and "why" in q_lower):
                last_run = db.query(RecommendationRun).filter(RecommendationRun.user_id == current_user.id).order_by(RecommendationRun.created_at.desc()).first()
                if last_run:
                    recs_data: List[Dict[str, Any]] = []
                    try:
                        recs_raw = last_run.recommendations_json
                        if isinstance(recs_raw, str):
                            import json as _json
                            recs_data = _json.loads(recs_raw)
                        else:
                            recs_data = recs_raw or []
                    except Exception:
                        recs_data = []
                    if recs_data:
                        import re as _re
                        m = _re.search(r"why(?:\s+was)?\s+([\w \-]+?)\s+recommended\s+over\s+([\w \-]+)", q_lower)
                        mat_a = mat_b = None
                        if m:
                            mat_a = m.group(1).strip()
                            mat_b = m.group(2).strip()
                        if not mat_a or not mat_b:
                            if len(recs_data) >= 2:
                                mat_a = str(recs_data[0].get("materialName") or recs_data[0].get("MaterialName") or recs_data[0].get("MaterialType") or "").lower()
                                mat_b = str(recs_data[1].get("materialName") or recs_data[1].get("MaterialName") or recs_data[1].get("MaterialType") or "").lower()
                            else:
                                mat_a = str(recs_data[0].get("materialName") or recs_data[0].get("MaterialName") or recs_data[0].get("MaterialType") or "").lower()
                                mat_b = None
                        name_map = {}
                        for rec in recs_data:
                            name = str(rec.get("materialName") or rec.get("MaterialName") or rec.get("MaterialType") or "").strip().lower()
                            if name:
                                name_map[name] = rec
                        info_a = name_map.get(mat_a) if mat_a else None
                        info_b = name_map.get(mat_b) if mat_b else None
                        if info_a and info_b:
                            score_a = float(info_a.get("rankingScore") or info_a.get("suitabilityScore") or 0.0)
                            score_b = float(info_b.get("rankingScore") or info_b.get("suitabilityScore") or 0.0)
                            cost_a = float(info_a.get("predictedCost") or info_a.get("predictedCostUSD") or info_a.get("predicted_cost_unit_usd") or 0.0)
                            cost_b = float(info_b.get("predictedCost") or info_b.get("predictedCostUSD") or info_b.get("predicted_cost_unit_usd") or 0.0)
                            co2_a = float(info_a.get("predictedCO2") or info_a.get("predictedCO2KG") or info_a.get("predicted_co2_unit_kg") or 0.0)
                            co2_b = float(info_b.get("predictedCO2") or info_b.get("predictedCO2KG") or info_b.get("predicted_co2_unit_kg") or 0.0)
                            reason_a = info_a.get("reason") or info_a.get("recommendationReason") or ""
                            reason_b = info_b.get("reason") or info_b.get("recommendationReason") or ""
                            comparison = []
                            comparison.append(f"In your last run, {mat_a.title()} had a higher ranking score ({score_a:.1f}) than {mat_b.title()} ({score_b:.1f}).")
                            if cost_a != cost_b:
                                comparison.append(f"It was predicted to cost {cost_a:.3f} USD/unit versus {cost_b:.3f} USD/unit.")
                            if co2_a != co2_b:
                                comparison.append(f"Its CO₂ impact was {co2_a:.3f} kg/unit compared with {co2_b:.3f} kg/unit.")
                            if reason_a:
                                comparison.append(f"Reasons for {mat_a.title()}: {reason_a}.")
                            if reason_b:
                                comparison.append(f"Reasons for {mat_b.title()}: {reason_b}.")
                            answer = " ".join(comparison)
                            answered = True
                        elif info_a:
                            answer = f"In your last run, {mat_a.title()} was recommended because: {info_a.get('reason') or info_a.get('recommendationReason') or 'it best met the sustainability and cost criteria.'}"
                            answered = True
                        else:
                            answer = "I could not find those materials in your last run. Please check the material names."
                            answered = True
                    else:
                        answer = "I couldn't retrieve your last recommendation list."
                        answered = True
                else:
                    answer = "You have no previous recommendation runs."
                    answered = True
        except Exception:
            answered = False

        if not answered:
            try:
                answer = chatbot.respond(payload.question, history=history_for_bot)
            except Exception:
                answer = chatbot.respond(payload.question, history=history)

        assistant_msg = ChatMessage(user_id=current_user.id, role="assistant", content=answer)
        db.add(assistant_msg)
        db.commit()

        try:
            log_audit(db, current_user, "chat", details={
                "question": payload.question,
                "answerPreview": answer[:100],
            })
        except Exception:
            pass

        return JSONResponse(content={"answer": answer})
    except Exception as e:
        logger.error("Chat error: %s", e)
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

# ---------- Materials and Comparison Endpoints ----------

@app.get("/materials-names")
def get_material_names(current_user: AppUser = Depends(get_current_user)) -> JSONResponse:
    try:
        prediction.load_models()
    except Exception as e:
        logger.warning("Unable to load models/materials: %s", e)
    names: List[str] = []
    try:
        df = prediction._materials_raw  # type: ignore[attr-defined]
        if df is not None:
            col = None
            if "MaterialName" in df.columns:
                col = "MaterialName"
            elif "MaterialType" in df.columns:
                col = "MaterialType"
            if col:
                names = sorted({str(x).strip() for x in df[col].dropna().unique() if str(x).strip()})
    except Exception as e:
        logger.warning("Failed to fetch material names: %s", e)
    return JSONResponse(content={"materials": names})


@app.post("/compare-materials")
def compare_materials(
    payload: CompareMaterialsRequest,
    current_user: AppUser = Depends(get_current_user),
) -> JSONResponse:
    """
    Compare a user-specified set of materials by predicting cost and CO₂ impact.
    IMPORTANT: This logic manually uses the prediction module's internal models/data
    to ensure we get results for ALL requested materials, not just the top 5
    returned by prediction.recommend().
    """
    try:
        # Ensure models and data are loaded
        prediction.load_models()
        # Access protected members to perform direct inference
        df_raw = getattr(prediction, "_materials_raw", None)
        cost_model = getattr(prediction, "_cost_model", None)
        co2_model = getattr(prediction, "_co2_model", None)
        
        if df_raw is None or cost_model is None or co2_model is None:
            raise RuntimeError("Models or materials data not initialized.")

        # Prepare product dict
        product = {
            "productName": payload.productName or "Generic Product",
            "category": payload.category,
            "weightKg": payload.weightKg,
            "fragility": payload.fragility,
            "maxBudget": payload.maxBudget,
            "shippingDistance": payload.shippingDistance,
            "moistureReq": payload.moistureReq,
            "oxygenSensitivity": payload.oxygenSensitivity,
            "preferredBiodegradable": payload.preferredBiodegradable,
            "preferredRecyclable": payload.preferredRecyclable,
        }

        # Identify requested materials (case-insensitive matching)
        requested_lower = {m.strip().lower(): m.strip() for m in payload.materials if m.strip()}
        
        # Filter raw dataframe to ONLY the requested materials
        # Try finding the material name column
        name_col = "MaterialName" if "MaterialName" in df_raw.columns else "MaterialType"
        
        # Create a lowercase mask for filtering
        # We process the dataframe to find matches
        mask = df_raw[name_col].astype(str).str.strip().str.lower().isin(requested_lower.keys())
        df_filtered = df_raw[mask].copy()
        
        results: List[Dict[str, Any]] = []
        found_materials_lower = set()

        # If we have matches, predict for them
        if not df_filtered.empty:
            # Build inference frame using internal helper
            # We access the protected helper function from prediction module
            X = prediction._build_inference_frame(product, df_filtered)
            
            # Predict
            pred_costs = cost_model.predict(X)
            pred_co2s = co2_model.predict(X)
            
            # Map back to results
            # Iterate through the filtered df to pair names with predictions
            df_filtered["_pred_cost"] = pred_costs
            df_filtered["_pred_co2"] = pred_co2s
            
            for _, row in df_filtered.iterrows():
                mat_name = str(row[name_col]).strip()
                mat_lower = mat_name.lower()
                
                # If this material was requested, add it to results
                if mat_lower in requested_lower:
                    found_materials_lower.add(mat_lower)
                    # We also need to determine an 'alternative'.
                    # For simplicity in this direct-mode, we won't compute a full recommendation 
                    # list just to find an alternative unless necessary. 
                    # If needed, we could run a full recommend() separate call, but that's expensive.
                    # We will leave alternativeMaterial as None for now or logic can be added later.
                    
                    results.append({
                        "material": mat_name, # Use DB name casing
                        "predictedCost": float(row["_pred_cost"]),
                        "predictedCO2": float(row["_pred_co2"]),
                        "alternativeMaterial": None 
                    })

        # Handle unknown materials (those requested but not found in DB)
        unknown_lower = [k for k in requested_lower.keys() if k not in found_materials_lower]
        
        if unknown_lower:
            # Reconstruct original casing for the prompt
            unknown_original = [requested_lower[k] for k in unknown_lower]
            try:
                question = (
                    "Estimate the predicted cost per unit (USD) and predicted CO₂ per unit (kg) for the following materials: "
                    + ", ".join(unknown_original)
                    + ". The product parameters are: category = "
                    + payload.category
                    + ", weight = "
                    + str(payload.weightKg)
                    + " kg, fragility = "
                    + str(payload.fragility)
                    + ", max budget = "
                    + str(payload.maxBudget)
                    + " USD/unit, shipping distance = "
                    + str(payload.shippingDistance)
                    + " km. Return the result strictly as a JSON array where each element has 'material', 'predictedCost' and 'predictedCO2' keys. "
                    + "Do not include any markdown formatting or backticks."
                )
                
                # Use chatbot's OpenAI integration
                answer = chatbot._openai_chat(question, history=[])
                
                if answer:
                    # Clean markdown if present (robustness fix)
                    clean_answer = answer.strip()
                    if clean_answer.startswith("```"):
                        clean_answer = clean_answer.strip("`").replace("json", "", 1).strip()
                    
                    import json as _json
                    try:
                        parsed = _json.loads(clean_answer)
                        if isinstance(parsed, list):
                            for item in parsed:
                                m_name = item.get("material") or item.get("name")
                                if not m_name:
                                    continue
                                pcost = item.get("predictedCost") or item.get("cost")
                                pco2 = item.get("predictedCO2") or item.get("co2")
                                results.append({
                                    "material": str(m_name),
                                    "predictedCost": float(pcost) if pcost is not None else None,
                                    "predictedCO2": float(pco2) if pco2 is not None else None,
                                    "alternativeMaterial": None,
                                })
                        else:
                            # Handle dict response
                            for mname, vals in (parsed.items() if isinstance(parsed, dict) else []):
                                pcost = vals.get("predictedCost") if isinstance(vals, dict) else None
                                pco2 = vals.get("predictedCO2") if isinstance(vals, dict) else None
                                results.append({
                                    "material": str(mname),
                                    "predictedCost": float(pcost) if pcost is not None else None,
                                    "predictedCO2": float(pco2) if pco2 is not None else None,
                                    "alternativeMaterial": None,
                                })
                    except Exception:
                        # Fallback regex parser
                        lines = clean_answer.split("\n")
                        import re as _re
                        for line in lines:
                            m = _re.match(r"\s*([\w \-]+)\s*[:\-]?\s*(?:cost|Cost)?\s*(\d+(?:\.\d+)?)\s*(?:USD)?\s*(?:[^\d]+)?(\d+(?:\.\d+)?)?", line)
                            if m:
                                name = m.group(1).strip()
                                cost_val = m.group(2)
                                co2_val = m.group(3)
                                results.append({
                                    "material": name,
                                    "predictedCost": float(cost_val) if cost_val else None,
                                    "predictedCO2": float(co2_val) if co2_val else None,
                                    "alternativeMaterial": None,
                                })
            except Exception as e:
                logger.warning("OpenAI estimation failed: %s", e)
                # Fill Nones for unknowns if AI fails
                for k in unknown_lower:
                    results.append({
                        "material": requested_lower[k],
                        "predictedCost": None,
                        "predictedCO2": None,
                        "alternativeMaterial": None,
                    })

        return JSONResponse(content={"results": results})

    except Exception as e:
        logger.error("Error in compare_materials: %s", e)
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))